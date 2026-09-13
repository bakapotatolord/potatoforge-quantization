"""Specialized worksheet renderers for analysis workbooks."""

from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from .analysis import (
    ANALYSIS_METHODS,
    ProfileReport,
    _measurement_value,
    _set_column_widths,
    _style_header,
    _style_table,
)
from .weight_audit import WeightAuditDocument


def write_tradeoffs_sheet(
    sheet: Worksheet,
    reports: tuple[ProfileReport, ...],
) -> None:
    sheet["A1"] = "Storage vs P95 tensor relative L2"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Lower-left is better: smaller output with lower error."
    sheet["A3"] = (
        "Each dot is a precomputed profile; the red dot marks the profile selected in Recommendations."
    )
    sheet["A3"].alignment = Alignment(wrap_text=True)

    header_row = 24
    headers = (
        "Profile",
        "Requested Target (GiB)",
        "Estimated Output (GiB)",
        "P95 Tensor Relative L2",
        "Selected?",
        "",
        "Selected Output",
        "Selected P95",
    )
    for column, value in enumerate(headers, start=1):
        sheet.cell(header_row, column, value)
    _style_header(sheet, header_row, len(headers))
    for index in range(len(reports)):
        row = header_row + 1 + index
        profile_row = index + 2
        sheet.cell(row, 1, f"='Profiles'!A{profile_row}")
        sheet.cell(row, 2, f"='Profiles'!B{profile_row}")
        sheet.cell(row, 3, f"='Profiles'!C{profile_row}")
        sheet.cell(row, 4, f"='Profiles'!F{profile_row}")
        sheet.cell(row, 5, f'=IF(A{row}=\'Recommendations\'!$B$2,"Selected","")')
        sheet.cell(row, 7, f'=IF(E{row}="Selected",C{row},"")')
        sheet.cell(row, 8, f'=IF(E{row}="Selected",D{row},"")')

    last_row = header_row + len(reports)
    _style_table(sheet, header_row, max(last_row, header_row), 5, f"A{header_row + 1}")
    for row in range(header_row + 1, last_row + 1):
        for column in (2, 3, 7):
            sheet.cell(row, column).number_format = "0.00"
        for column in (4, 8):
            sheet.cell(row, column).number_format = "0.000000"
    sheet.column_dimensions["F"].hidden = True
    sheet.column_dimensions["G"].hidden = True
    sheet.column_dimensions["H"].hidden = True
    _set_column_widths(sheet, 8)

    if not reports:
        return
    chart = ScatterChart()
    chart.title = "Output size vs P95 tensor relative L2"
    chart.style = 13
    chart.height = 10
    chart.width = 18
    chart.x_axis.title = "Estimated output size (GiB)"
    chart.y_axis.title = "P95 tensor relative L2"
    chart.x_axis.numFmt = "0.0"
    chart.y_axis.numFmt = "0.0%"

    x_values = Reference(sheet, min_col=3, min_row=header_row + 1, max_row=last_row)
    y_values = Reference(sheet, min_col=4, min_row=header_row + 1, max_row=last_row)
    all_profiles = Series(y_values, x_values, title="All profiles")
    all_profiles.marker.symbol = "circle"
    all_profiles.marker.size = 6
    all_profiles.graphicalProperties.line.noFill = True
    all_profiles.marker.graphicalProperties.solidFill = "5B9BD5"
    all_profiles.marker.graphicalProperties.line.solidFill = "5B9BD5"
    chart.series.append(all_profiles)

    selected_x = Reference(
        sheet,
        min_col=7,
        min_row=header_row + 1,
        max_row=last_row,
    )
    selected_y = Reference(
        sheet,
        min_col=8,
        min_row=header_row + 1,
        max_row=last_row,
    )
    selected_profile = Series(selected_y, selected_x, title="Selected")
    selected_profile.marker.symbol = "circle"
    selected_profile.marker.size = 10
    selected_profile.graphicalProperties.line.noFill = True
    selected_profile.marker.graphicalProperties.solidFill = "C00000"
    selected_profile.marker.graphicalProperties.line.solidFill = "C00000"
    chart.series.append(selected_profile)
    sheet.add_chart(chart, "A4")


def write_errors_sheet(
    sheet: Worksheet,
    audit: WeightAuditDocument,
) -> None:
    headers = ["Tensor", *(label for _, label in ANALYSIS_METHODS)]
    sheet.append(headers)
    for result in audit["results"]:
        row: list[object] = [
            result["tensor_name"],
            *(
                _measurement_value(result, method)
                for method, _ in ANALYSIS_METHODS
            ),
        ]
        sheet.append(row)

    last_row = 1 + len(audit["results"])
    method_last_column = 1 + len(ANALYSIS_METHODS)
    last_column = len(headers)
    _style_table(sheet, 1, last_row, last_column, "B2")
    for row in range(2, last_row + 1):
        columns = range(2, method_last_column + 1)
        values = [sheet.cell(row, column).value for column in columns]
        least_error = min(
            (value for value in values if value is not None),
            default=None,
        )
        for column, value in zip(columns, values):
            sheet.cell(row, column).number_format = "0.000000"
            if least_error is not None and value == least_error:
                sheet.cell(row, column).font = Font(bold=True)
    _set_column_widths(sheet, last_column)
