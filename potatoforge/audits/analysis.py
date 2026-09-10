"""Render existing weight-audit measurements as a workbook or text report."""

from collections import Counter
from math import fsum, prod
from pathlib import Path
from typing import Final, NamedTuple

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..headers.source_header import SourceModelHeader
from ..planning import source_bytes
from ..profiles import resolve_profile
from .profile_optimizer import (
    OptimizedProfile,
    SUPPORTED_METHODS,
    audited_global_relative_l2,
    method_for_action,
    next_method_choice,
)
from .weight_audit import (
    QuantizationMethod,
    WeightAuditDocument,
    WeightAuditRecord,
)


ANALYSIS_METHODS: Final[tuple[tuple[QuantizationMethod, str], ...]] = (
    ("convrot_w4a4", "W4A4 ConvRot"),
    ("convrot_w4a4_mse", "W4A4 ConvRot + MSE"),
    ("int6", "INT6"),
    ("int6_convrot", "INT6 ConvRot"),
    ("int8", "INT8"),
    ("int8_convrot", "INT8 ConvRot"),
)
PROFILE_METHODS: Final[tuple[tuple[QuantizationMethod, str], ...]] = (
    ("bf16", "BF16"),
    *ANALYSIS_METHODS,
)

_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
_GOOD_BAD_SCALE = ColorScaleRule(
    start_type="min",
    start_color="63BE7B",
    mid_type="percentile",
    mid_value=50,
    mid_color="FFEB84",
    end_type="max",
    end_color="F8696B",
)

_KIB = 1024
_MIB = 1024**2
_GIB = 1024**3


class ProfileTensor(NamedTuple):
    tensor_name: str
    method: QuantizationMethod
    storage_mib: float | None
    error: float | None
    reconstruction_sse: float | None
    upgrade_method: QuantizationMethod | None
    extra_storage_mib: float | None
    upgrade_sse_reduction: float | None
    sse_reduction_per_mib: float | None
    energy_fraction: float | None
    normalized_sse: float | None


class ProfileReport(NamedTuple):
    label: str
    optimized: OptimizedProfile
    mean_error: float | None
    p95_error: float | None
    max_error: float | None
    audited_global_relative_l2: float | None
    normalized_sse: float | None
    tensors_under_1_percent: int
    tensors_under_5_percent: int
    method_counts: Counter[QuantizationMethod]
    tensor_rows: tuple[ProfileTensor, ...]


def format_bytes(byte_count: int) -> str:
    for unit, divisor in (("GiB", _GIB), ("MiB", _MIB), ("KiB", _KIB)):
        if byte_count >= divisor:
            return f"{byte_count / divisor:.2f} {unit}"
    return f"{byte_count} B"


def _format_shape(shape: list[int]) -> str:
    return " x ".join(str(dimension) for dimension in shape)


def _source_file_bytes(audit: WeightAuditDocument) -> int:
    return Path(audit["source_path"]).stat().st_size


def _source_payload_bytes(source_header: SourceModelHeader) -> int:
    return sum(source_bytes(descriptor) for descriptor in source_header.tensors.values())


def _find_result(
    audit: WeightAuditDocument,
    tensor_name: str,
) -> WeightAuditRecord:
    for result in audit["results"]:
        if result["tensor_name"] == tensor_name:
            return result
    raise ValueError(f"Tensor not found in weight audit: {tensor_name}")


def _measurement_value(
    result: WeightAuditRecord,
    method: QuantizationMethod,
) -> float | None:
    measurement = result["methods"].get(method)
    return None if measurement is None else measurement["relative_l2_error"]


def _tensor_row(
    result: WeightAuditRecord,
    source_header: SourceModelHeader,
    total_source_bytes: int,
) -> tuple[str, str, int, int, float | None]:
    tensor_name = result["tensor_name"]
    descriptor = source_header.tensors[tensor_name]
    tensor_source_bytes = source_bytes(descriptor)
    model_share = (
        None
        if total_source_bytes == 0
        else tensor_source_bytes / total_source_bytes
    )
    return (
        tensor_name,
        _format_shape(result["shape"]),
        prod(result["shape"]),
        tensor_source_bytes,
        model_share,
    )


def _style_table(
    sheet: Worksheet,
    header_row: int,
    last_row: int,
    last_column: int,
    freeze_panes: str,
) -> None:
    for cell in sheet[header_row]:
        if cell.column <= last_column:
            cell.font = Font(bold=True)
    sheet.freeze_panes = freeze_panes
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(last_column)}{max(last_row, header_row)}"
    )


def _set_column_widths(sheet: Worksheet, last_column: int) -> None:
    for column_index in range(1, last_column + 1):
        values = (
            len(str(cell.value))
            for column_cells in sheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
            )
            for cell in column_cells
            if cell.value is not None
        )
        width = max(values, default=10) + 2
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(width, 12),
            60,
        )


def _percentile(values: list[float], percentile: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * percentile
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    fraction = position - lower
    return ordered[lower] + (ordered[upper] - ordered[lower]) * fraction


def _is_excluded(
    tensor_name: str,
    excluded_prefixes: tuple[str, ...],
    excluded_suffixes: tuple[str, ...],
) -> bool:
    return tensor_name.startswith(excluded_prefixes) or tensor_name.endswith(
        excluded_suffixes
    )


def _profile_label(
    index: int,
    profile_count: int,
    optimized: OptimizedProfile,
) -> str:
    size = optimized.output_bytes / _GIB
    if index == profile_count - 1 and not optimized.profile["rules"]:
        return f"All BF16 ({size:.2f} GiB)"
    if index == 0:
        return f"Minimum ({size:.2f} GiB)"
    return f"Target {optimized.target_bytes / _GIB:.2f} GiB"


def _build_profile_report(
    audit: WeightAuditDocument,
    optimized: OptimizedProfile,
    label: str,
    allowed_methods: frozenset[QuantizationMethod],
    excluded_prefixes: tuple[str, ...],
    excluded_suffixes: tuple[str, ...],
) -> ProfileReport:
    rows: list[ProfileTensor] = []
    errors: list[float] = []
    method_counts: Counter[QuantizationMethod] = Counter()
    total_weight_energy = fsum(
        float(result["weight_l2_sq"]) for result in audit["results"]
    )
    for result in audit["results"]:
        tensor_name = result["tensor_name"]
        method = method_for_action(
            resolve_profile(optimized.profile, tensor_name)
        )
        measurement = result["methods"][method]
        storage_bytes = measurement["storage_bytes"]
        error = measurement["relative_l2_error"]
        reconstruction_sse = (
            None
            if error is None
            else float(error) ** 2 * result["weight_l2_sq"]
        )
        energy_fraction = (
            None
            if total_weight_energy <= 0
            else result["weight_l2_sq"] / total_weight_energy
        )
        normalized_sse = (
            None
            if reconstruction_sse is None or total_weight_energy <= 0
            else reconstruction_sse / total_weight_energy
        )
        if error is not None:
            errors.append(float(error))
        method_counts[method] += 1

        tensor_methods = (
            frozenset(("bf16",))
            if _is_excluded(
                tensor_name,
                excluded_prefixes,
                excluded_suffixes,
            )
            else allowed_methods
        )
        upgrade = next_method_choice(result, method, tensor_methods)
        if upgrade is None or storage_bytes is None or error is None:
            rows.append(
                ProfileTensor(
                    tensor_name,
                    method,
                    None if storage_bytes is None else storage_bytes / _MIB,
                    error,
                    reconstruction_sse,
                    None,
                    None,
                    None,
                    None,
                    energy_fraction,
                    normalized_sse,
                )
            )
            continue

        extra_storage_bytes = upgrade.storage_bytes - storage_bytes
        assert reconstruction_sse is not None
        upgrade_sse_reduction = reconstruction_sse - upgrade.reconstruction_sse
        sse_reduction_per_mib = (
            None
            if extra_storage_bytes <= 0
            else upgrade_sse_reduction / extra_storage_bytes * _MIB
        )
        rows.append(
            ProfileTensor(
                tensor_name,
                method,
                storage_bytes / _MIB,
                float(error),
                reconstruction_sse,
                upgrade.method,
                extra_storage_bytes / _MIB,
                upgrade_sse_reduction,
                sse_reduction_per_mib,
                energy_fraction,
                normalized_sse,
            )
        )

    return ProfileReport(
        label=label,
        optimized=optimized,
        mean_error=None if not errors else sum(errors) / len(errors),
        p95_error=_percentile(errors, 0.95),
        max_error=max(errors, default=None),
        audited_global_relative_l2=audited_global_relative_l2(
            audit,
            optimized.reconstruction_sse,
        ),
        normalized_sse=(
            None
            if total_weight_energy <= 0
            else optimized.reconstruction_sse / total_weight_energy
        ),
        tensors_under_1_percent=sum(error <= 0.01 for error in errors),
        tensors_under_5_percent=sum(error <= 0.05 for error in errors),
        method_counts=method_counts,
        tensor_rows=tuple(rows),
    )


def _build_profile_reports(
    audit: WeightAuditDocument,
    optimized_profiles: tuple[OptimizedProfile, ...],
    allowed_methods: frozenset[QuantizationMethod],
    excluded_prefixes: tuple[str, ...],
    excluded_suffixes: tuple[str, ...],
) -> tuple[ProfileReport, ...]:
    profile_count = len(optimized_profiles)
    return tuple(
        _build_profile_report(
            audit,
            optimized,
            _profile_label(index, profile_count, optimized),
            allowed_methods,
            excluded_prefixes,
            excluded_suffixes,
        )
        for index, optimized in enumerate(optimized_profiles)
    )


def _style_header(sheet: Worksheet, row: int, last_column: int) -> None:
    for column in range(1, last_column + 1):
        cell = sheet.cell(row, column)
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _profile_lookup_formula(
    return_column: str,
    key_reference: str,
    last_row: int,
) -> str:
    return (
        f'=IFERROR(INDEX(\'Profile Data\'!${return_column}$2:'
        f'${return_column}${last_row},MATCH({key_reference},'
        f'\'Profile Data\'!$A$2:$A${last_row},0)),"N/A")'
    )


def _profile_summary_formula(return_column: str, last_row: int) -> str:
    return (
        f'=IFERROR(INDEX(\'Profiles\'!${return_column}$2:'
        f'${return_column}${last_row},MATCH($B$2,'
        f'\'Profiles\'!$A$2:$A${last_row},0)),"N/A")'
    )


def _write_summary(
    sheet: Worksheet,
    audit: WeightAuditDocument,
    source_header: SourceModelHeader,
) -> None:
    source_file_bytes = _source_file_bytes(audit)
    total_source_bytes = _source_payload_bytes(source_header)
    results = audit["results"]

    sheet.append(["Field", "Value"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for label, value in (
        ("Source", Path(audit["source_path"]).name),
        ("Original model size (MiB)", source_file_bytes / _MIB),
        ("Total source tensor count", len(source_header.tensors)),
        ("Audited weight count", len(results)),
        (
            "Total audited parameters",
            sum(prod(result["shape"]) for result in results),
        ),
        ("Audit format version", audit["format_version"]),
    ):
        sheet.append([label, value])

    header_row = 9
    sheet.cell(header_row, 1, "Tensor")
    sheet.cell(header_row, 2, "Shape")
    sheet.cell(header_row, 3, "Elements")
    sheet.cell(header_row, 4, "Original Size MiB")
    sheet.cell(header_row, 5, "Model Share %")

    for result in results:
        name, shape, elements, tensor_bytes, model_share = _tensor_row(
            result,
            source_header,
            total_source_bytes,
        )
        sheet.append(
            [
                name,
                shape,
                elements,
                tensor_bytes / _MIB,
                model_share,
            ]
        )

    last_row = header_row + len(results)
    _style_table(sheet, header_row, last_row, 5, f"A{header_row + 1}")
    for row in range(header_row + 1, last_row + 1):
        sheet.cell(row, 4).number_format = "0.00"
        sheet.cell(row, 5).number_format = "0.00%"
    sheet.cell(2, 2).number_format = "0.00"
    _set_column_widths(sheet, 5)


def _report_value(value: float | None) -> float | str:
    return "N/A" if value is None else value


def _write_profiles(
    sheet: Worksheet,
    audit: WeightAuditDocument,
    reports: tuple[ProfileReport, ...],
) -> None:
    source_file_bytes = _source_file_bytes(audit)
    headers = [
        "Profile",
        "Requested Target (GiB)",
        "Estimated Output (GiB)",
        "Space Saved %",
        "Mean Tensor Relative L2",
        "P95 Tensor Relative L2",
        "Maximum Tensor Relative L2",
        "Tensors <=1% Relative L2",
        "Tensors <=5% Relative L2",
        "Reconstruction SSE",
        "Audited Global Relative L2",
        "Normalized Reconstruction SSE",
        *(label for _, label in PROFILE_METHODS),
    ]
    sheet.append(headers)
    _style_header(sheet, 1, len(headers))
    for report in reports:
        optimized = report.optimized
        output_gib = optimized.output_bytes / _GIB
        saved_fraction = (
            None
            if source_file_bytes == 0
            else 1 - optimized.output_bytes / source_file_bytes
        )
        sheet.append(
            [
                report.label,
                optimized.target_bytes / _GIB,
                output_gib,
                saved_fraction,
                _report_value(report.mean_error),
                _report_value(report.p95_error),
                _report_value(report.max_error),
                report.tensors_under_1_percent,
                report.tensors_under_5_percent,
                optimized.reconstruction_sse,
                _report_value(report.audited_global_relative_l2),
                _report_value(report.normalized_sse),
                *(
                    report.method_counts.get(method, 0)
                    for method, _ in PROFILE_METHODS
                ),
            ]
        )

    last_row = max(1, 1 + len(reports))
    _style_table(sheet, 1, last_row, len(headers), "A2")
    for row in range(2, last_row + 1):
        for column in (2, 3):
            sheet.cell(row, column).number_format = "0.00"
        sheet.cell(row, 4).number_format = "0.0%"
        for column in (5, 6, 7):
            sheet.cell(row, column).number_format = "0.000000"
        for column in (10, 11, 12):
            sheet.cell(row, column).number_format = "0.000000"
    sheet.row_dimensions[1].height = 30
    _set_column_widths(sheet, len(headers))


def _write_profile_data(
    sheet: Worksheet,
    reports: tuple[ProfileReport, ...],
) -> int:
    headers = [
        "Lookup Key",
        "Profile",
        "Tensor",
        "Recommended Method",
        "Storage MiB",
        "Relative L2 Error",
        "Reconstruction SSE",
        "Upgrade Method",
        "Extra Storage MiB",
        "Upgrade SSE Reduction",
        "SSE Reduction / MiB",
        "Energy Fraction",
        "Normalized Reconstruction SSE",
    ]
    sheet.append(headers)
    _style_header(sheet, 1, len(headers))
    for report in reports:
        for tensor in report.tensor_rows:
            sheet.append(
                [
                    f"{report.label}|{tensor.tensor_name}",
                    report.label,
                    tensor.tensor_name,
                    tensor.method,
                    _report_value(tensor.storage_mib),
                    _report_value(tensor.error),
                    _report_value(tensor.reconstruction_sse),
                    tensor.upgrade_method or "N/A",
                    _report_value(tensor.extra_storage_mib),
                    _report_value(tensor.upgrade_sse_reduction),
                    _report_value(tensor.sse_reduction_per_mib),
                    _report_value(tensor.energy_fraction),
                    _report_value(tensor.normalized_sse),
                ]
            )

    last_row = max(1, sheet.max_row)
    _style_table(sheet, 1, last_row, len(headers), "B2")
    sheet.column_dimensions["A"].hidden = True
    for row in range(2, last_row + 1):
        for column in (5, 9):
            sheet.cell(row, column).number_format = "0.00"
        for column in (6, 7, 11, 13):
            sheet.cell(row, column).number_format = "0.000000"
        sheet.cell(row, 12).number_format = "0.00%"
    _set_column_widths(sheet, len(headers))
    return last_row


def _write_recommendations(
    sheet: Worksheet,
    audit: WeightAuditDocument,
    reports: tuple[ProfileReport, ...],
    selected_label: str | None,
    profile_data_last_row: int,
) -> None:
    sheet.append(["Field", "Value"])
    _style_header(sheet, 1, 2)
    if not reports or selected_label is None:
        sheet.append(
            [
                "Status",
                "No profile sweep was generated. Recommendations were not generated.",
            ]
        )
        _set_column_widths(sheet, 2)
        return

    sheet.append(["Selected Profile", selected_label])
    for label, column in (
        ("Requested Target (GiB)", "B"),
        ("Estimated Output (GiB)", "C"),
        ("Space Saved", "D"),
        ("Mean Tensor Relative L2", "E"),
        ("P95 Tensor Relative L2", "F"),
        ("Maximum Tensor Relative L2", "G"),
        ("Tensors <=1% Relative L2", "H"),
        ("Tensors <=5% Relative L2", "I"),
        ("Reconstruction SSE", "J"),
        ("Audited Global Relative L2", "K"),
        ("Normalized Reconstruction SSE", "L"),
    ):
        sheet.append([label, _profile_summary_formula(column, len(reports) + 1)])
    sheet.append(
        [
            "Warning",
            "Audit relative L2 error is a weight-reconstruction proxy; it is not runtime or image-quality validation.",
        ]
    )

    header_row = 15
    headers = (
        "Tensor",
        "Recommended Method",
        "Storage MiB",
        "Relative L2 Error",
        "Reconstruction SSE",
        "Upgrade Method",
        "Extra Storage MiB",
        "Upgrade SSE Reduction",
        "SSE Reduction / MiB",
        "Energy Fraction",
        "Normalized Reconstruction SSE",
    )
    for column, value in enumerate(headers, start=1):
        sheet.cell(header_row, column, value)
    _style_header(sheet, header_row, len(headers))
    for index, result in enumerate(audit["results"], start=header_row + 1):
        tensor_name = result["tensor_name"]
        key_reference = f'$B$2&"|"&$A{index}'
        sheet.cell(index, 1, tensor_name)
        for column, return_column in enumerate(
            ("D", "E", "F", "G", "H", "I", "J", "K", "L", "M"),
            start=2,
        ):
            sheet.cell(
                index,
                column,
                _profile_lookup_formula(return_column, key_reference, profile_data_last_row),
            )

    last_row = header_row + len(audit["results"])
    _style_table(sheet, header_row, last_row, len(headers), f"A{header_row + 1}")
    for row in range(header_row + 1, last_row + 1):
        for column in (3, 7):
            sheet.cell(row, column).number_format = "0.00"
        for column in (4, 5, 8, 9, 11):
            sheet.cell(row, column).number_format = "0.000000"
        sheet.cell(row, 10).number_format = "0.00%"
    if last_row >= header_row + 1:
        sheet.conditional_formatting.add(
            f"D{header_row + 1}:D{last_row}",
            _GOOD_BAD_SCALE,
        )
        sheet.conditional_formatting.add(
            f"H{header_row + 1}:H{last_row}",
            _GOOD_BAD_SCALE,
        )
    for row, number_format in (
        (3, "0.00"),
        (4, "0.00"),
        (5, "0.0%"),
        (6, "0.000000"),
        (7, "0.000000"),
        (8, "0.000000"),
        (11, "0.000000"),
        (12, "0.000000"),
        (13, "0.000000"),
    ):
        sheet.cell(row, 2).number_format = number_format
    sheet.cell(14, 2).fill = _WARNING_FILL
    sheet.cell(14, 2).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells("B14:K14")
    sheet.row_dimensions[14].height = 34
    sheet.row_dimensions[header_row].height = 30
    for column, width in {
        "A": 52,
        "B": 24,
        "C": 14,
        "D": 18,
        "E": 22,
        "F": 18,
        "G": 18,
        "H": 22,
        "I": 22,
        "J": 16,
        "K": 24,
    }.items():
        sheet.column_dimensions[column].width = width


def write_analysis_workbook(
    audit: WeightAuditDocument,
    source_header: SourceModelHeader,
    output_path: str | Path,
    optimized: OptimizedProfile | None = None,
    overwrite: bool = False,
    profile_sweep: tuple[OptimizedProfile, ...] | None = None,
    allowed_methods: frozenset[QuantizationMethod] = SUPPORTED_METHODS,
    excluded_prefixes: tuple[str, ...] = (),
    excluded_suffixes: tuple[str, ...] = (),
) -> None:
    from .analysis_sheets import write_errors_sheet, write_tradeoffs_sheet

    output = Path(output_path)
    if Path(audit["source_path"]).resolve() == output.resolve():
        raise ValueError("Source and output paths cannot be the same.")
    if output.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    recommendations = workbook.create_sheet("Recommendations")
    errors = workbook.create_sheet("Errors")
    tradeoffs = workbook.create_sheet("Trade-offs")
    profiles = workbook.create_sheet("Profiles")
    profile_data = workbook.create_sheet("Profile Data")

    optimized_profiles = (
        profile_sweep
        if profile_sweep is not None
        else () if optimized is None else (optimized,)
    )
    reports = _build_profile_reports(
        audit,
        optimized_profiles,
        allowed_methods,
        excluded_prefixes,
        excluded_suffixes,
    )
    selected = optimized
    if selected is None and reports:
        selected = reports[-1].optimized
    selected_label = next(
        (report.label for report in reports if report.optimized == selected),
        None,
    )

    _write_summary(summary, audit, source_header)
    write_errors_sheet(errors, audit)
    _write_profiles(profiles, audit, reports)
    profile_data_last_row = _write_profile_data(profile_data, reports)
    _write_recommendations(
        recommendations,
        audit,
        reports,
        selected_label,
        profile_data_last_row,
    )
    write_tradeoffs_sheet(tradeoffs, reports)

    if reports:
        profiles_last_row = len(reports) + 1
        workbook.defined_names.add(
            DefinedName(
                "ProfileOptions",
                attr_text=f"'Profiles'!$A$2:$A${profiles_last_row}",
            )
        )
        validation = DataValidation(
            type="list",
            formula1="=ProfileOptions",
            allow_blank=False,
        )
        validation.error = "Choose one of the generated profile options."
        validation.errorTitle = "Invalid profile"
        validation.prompt = "Select a precomputed target profile."
        validation.promptTitle = "Profile selection"
        recommendations.add_data_validation(validation)
        validation.add(recommendations["B2"])

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)


def print_tensor_analysis(
    audit: WeightAuditDocument,
    source_header: SourceModelHeader,
    tensor_name: str,
    optimized: OptimizedProfile | None = None,
) -> None:
    result = _find_result(audit, tensor_name)
    total_source_bytes = _source_payload_bytes(source_header)
    _, shape, elements, tensor_bytes, model_share = _tensor_row(
        result,
        source_header,
        total_source_bytes,
    )

    print("Tensor")
    print("------")
    print(f"Name:          {tensor_name}")
    print(f"Shape:         {shape}")
    print(f"Elements:      {elements:,}")
    print(f"Original Size: {format_bytes(tensor_bytes)}")
    print(
        "Model Share:   "
        + ("n/a" if model_share is None else f"{model_share:.2%}")
    )
    print()
    print("Quantization Error")
    print("------------------")
    for method, label in ANALYSIS_METHODS:
        value = _measurement_value(result, method)
        display = "n/a" if value is None else f"{value:.6f}"
        print(f"{label:<28}{display}")

    activation = result.get("activation")
    if activation is not None:
        print()
        print("Activation calibration")
        print("----------------------")
        print(f"Metric:             {activation['metric_variant']}")
        print(
            "Baseline:           "
            + str(activation.get("calibration_baseline") or "n/a")
        )
        print(
            "Reference:          "
            + str(activation.get("activation_reference") or "n/a")
        )
        print(
            "Candidate:          "
            + str(activation.get("candidate_format") or "n/a")
        )
        print(f"Samples:            {activation['sample_count']}")
        print(f"Invocations:        {activation['invocation_count']}")
        print(f"Input features:     {activation['input_features']}")
        energy = activation["activation_energy_sum"]
        print(
            "Activation energy:  "
            + ("n/a" if energy is None else f"{energy:.6e}")
        )
        print()
        print("Activation-aware error")
        print("----------------------")
        error = activation["error"]
        print(
            "Raw score:          "
            + ("n/a" if error is None else f"{error:.6e}")
        )
        print(f"Status:              {activation['status']}")

    if optimized is None:
        return

    method = method_for_action(resolve_profile(optimized.profile, tensor_name))
    measurement = result["methods"][method]
    print()
    print("Recommendation")
    print("--------------")
    print(f"Method:          {method}")
    print(
        "Storage:         "
        + (
            "n/a"
            if measurement["storage_bytes"] is None
            else format_bytes(measurement["storage_bytes"])
        )
    )
    print(
        "Relative L2:     "
        + (
            "n/a"
            if measurement["relative_l2_error"] is None
            else f"{measurement['relative_l2_error']:.6f}"
        )
    )
    print(
        "Reconstruction SSE: "
        + (
            "n/a"
            if measurement["relative_l2_error"] is None
            else f"{float(measurement['relative_l2_error']) ** 2 * result['weight_l2_sq']:.6f}"
        )
    )


def print_activation_ranking(audit: WeightAuditDocument) -> None:
    if not isinstance(audit, dict):
        return
    status_counts: Counter[str] = Counter()
    metric_variants: set[str] = set()
    references: set[str] = set()
    baselines: set[str] = set()
    ranked: list[tuple[str, float]] = []
    for result in audit.get("results", ()):
        activation = result.get("activation")
        if activation is None:
            continue
        status_counts[activation["status"]] += 1
        metric_variants.add(activation["metric_variant"])
        reference = activation.get("activation_reference")
        if reference is not None:
            references.add(reference)
        baseline = activation.get("calibration_baseline")
        if baseline is not None:
            baselines.add(baseline)
        if activation["status"] == "ok" and activation["error"] is not None:
            ranked.append((result["tensor_name"], activation["error"]))

    if not status_counts:
        return

    print()
    print("Activation-aware ConvRot W4A4 ranking")
    print("--------------------------------------")
    if metric_variants:
        print(f"Metric: {', '.join(sorted(metric_variants))}")
    if baselines:
        print(f"Baseline: {', '.join(sorted(baselines))}")
    if references:
        print(f"Reference: {', '.join(sorted(references))}")
    for index, (tensor_name, error) in enumerate(
        sorted(ranked, key=lambda item: item[1], reverse=True)[:20],
        start=1,
    ):
        print(f"{index:>3}  {tensor_name:<52} {error:.3e}")
    if not ranked:
        print("No tensors with a valid activation score.")

    coverage = [f"scored: {len(ranked)}"]
    if status_counts["missing_calibration"]:
        coverage.append(
            f"missing: {status_counts['missing_calibration']}"
        )
    unavailable = sum(
        count
        for status, count in status_counts.items()
        if status not in {"ok", "missing_calibration"}
    )
    if unavailable:
        coverage.append(f"unavailable: {unavailable}")
    print("Coverage: " + ", ".join(coverage))
