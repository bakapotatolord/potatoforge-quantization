"""Compare every activation metric in one Excel workbook."""

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from ..calibration import ActivationCalibration, LayerCalibration
from ..headers.source_header import SourceModelHeader, read_source_model_header
from ..planning import source_bytes
from ..profiles import resolve_profile
from .activation_audit import (
    ACTIVATION_AUDIT_METRICS,
    ActivationAuditCache,
    score_activation_audit,
)
from .activation_profiles import (
    ActivationProfileResult,
    generate_activation_cache_profile,
)


ACTIVATION_METRICS: Final[tuple[str, ...]] = ACTIVATION_AUDIT_METRICS

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_ACTION_FILLS = {
    "keep": PatternFill("solid", fgColor="E7E6E6"),
    "convrot_w4a4": PatternFill("solid", fgColor="FFF2CC"),
    "int8_convrot": PatternFill("solid", fgColor="D9EAD3"),
}
_NUMBER_FORMAT = "0.000000000000"


def generate_activation_comparison_workbook(
    audit_cache: ActivationAuditCache | str | Path,
    activation_calibration: ActivationCalibration | str | Path,
    output_path: str | Path,
    *,
    source_path: str | Path | None = None,
    target_bytes: int | None = None,
    promotion_budget_bytes: int | None = None,
    allowed_methods: frozenset[str] | None = None,
    baseline_method: str = "bf16",
    excluded_prefixes: tuple[str, ...] = (),
    excluded_suffixes: tuple[str, ...] = (),
    metrics: Sequence[str] = ACTIVATION_METRICS,
    top_n: int = 20,
    overwrite: bool = False,
) -> dict[str, object]:
    """Score and optimize all requested metrics without writing profiles."""
    output = Path(output_path)
    if output.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    if top_n < 1:
        raise ValueError("top_n must be an integer >= 1.")

    selected_metrics = tuple(metrics)
    if not selected_metrics:
        raise ValueError("At least one activation metric is required.")
    if len(set(selected_metrics)) != len(selected_metrics):
        raise ValueError("Activation metrics must be unique.")

    calibration = (
        activation_calibration
        if isinstance(activation_calibration, ActivationCalibration)
        else ActivationCalibration.load(activation_calibration)
    )
    cache = (
        audit_cache
        if isinstance(audit_cache, ActivationAuditCache)
        else ActivationAuditCache.load(
            audit_cache,
            calibration=calibration,
        )
    )
    source = cache.source_model_path if source_path is None else source_path
    source_header = read_source_model_header(source)
    score_reports: dict[str, dict[str, object]] = {}
    profiles: dict[str, ActivationProfileResult] = {}
    for metric in selected_metrics:
        report = score_activation_audit(cache, calibration, metric=metric)
        score_reports[metric] = report
        profiles[metric] = generate_activation_cache_profile(
            cache,
            calibration,
            source_path=source,
            source_header=source_header,
            target_bytes=target_bytes,
            promotion_budget_bytes=promotion_budget_bytes,
            profile_id=f"activation-compare-{metric}",
            metric=metric,
            score_report=report,
            allowed_methods=allowed_methods,
            baseline_method=baseline_method,
            excluded_prefixes=excluded_prefixes,
            excluded_suffixes=excluded_suffixes,
        )

    _write_workbook(
        output,
        cache,
        calibration,
        source_header,
        score_reports,
        profiles,
        source_path=source,
        baseline_method=baseline_method,
        allowed_methods=allowed_methods,
        excluded_prefixes=excluded_prefixes,
        excluded_suffixes=excluded_suffixes,
        top_n=top_n,
    )
    return {
        "format": "potatoforge_activation_metric_comparison",
        "version": 1,
        "workbook_path": str(output.resolve()),
        "source_path": str(Path(source).resolve()),
        "calibration_session_id": cache.calibration_session_id,
        "metric_count": len(selected_metrics),
        "metrics": list(selected_metrics),
        "layer_count": len(cache.measurements),
        "candidate_count": sum(
            len(report["results"])
            for report in score_reports.values()
            if isinstance(report.get("results"), list)
        ),
    }


def _write_workbook(
    output: Path,
    cache: ActivationAuditCache,
    calibration: ActivationCalibration,
    source_header: SourceModelHeader,
    score_reports: Mapping[str, Mapping[str, object]],
    profiles: Mapping[str, ActivationProfileResult],
    *,
    source_path: str | Path,
    baseline_method: str,
    allowed_methods: frozenset[str] | None,
    excluded_prefixes: tuple[str, ...],
    excluded_suffixes: tuple[str, ...],
    top_n: int,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    candidate_sheet = workbook.create_sheet("Candidate Metrics")
    assignments_sheet = workbook.create_sheet("Assignments")
    comparison_sheet = workbook.create_sheet("Profile Comparison")
    rankings_sheet = workbook.create_sheet("Rankings")

    metrics = tuple(score_reports)
    candidate_methods = tuple(
        sorted(
            (
                set(cache.requested_methods)
                if allowed_methods is None
                else set(allowed_methods)
            )
            | ({baseline_method} if baseline_method != "bf16" else set())
        )
    )
    score_rows = {
        metric: _score_rows(report)
        for metric, report in score_reports.items()
    }
    assignment_rows = _assignment_rows(
        cache,
        source_header,
        profiles,
        score_rows,
        candidate_methods,
    )
    _write_summary(
        summary,
        cache,
        calibration,
        source_header,
        profiles,
        metrics,
        source_path=source_path,
        baseline_method=baseline_method,
        allowed_methods=allowed_methods,
        excluded_prefixes=excluded_prefixes,
        excluded_suffixes=excluded_suffixes,
    )
    _write_candidate_metrics(
        candidate_sheet,
        cache,
        calibration,
        score_rows,
        candidate_methods,
    )
    _write_table(
        assignments_sheet,
        _assignment_headers(candidate_methods),
        assignment_rows,
        "AssignmentsTable",
    )
    _write_profile_comparison(
        comparison_sheet,
        cache,
        profiles,
    )
    _write_rankings(
        rankings_sheet,
        cache,
        score_rows,
        candidate_methods,
        top_n,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def _write_summary(
    sheet: Worksheet,
    cache: ActivationAuditCache,
    calibration: ActivationCalibration,
    source_header: SourceModelHeader,
    profiles: Mapping[str, ActivationProfileResult],
    metrics: Sequence[str],
    *,
    source_path: str | Path,
    baseline_method: str,
    allowed_methods: frozenset[str] | None,
    excluded_prefixes: tuple[str, ...],
    excluded_suffixes: tuple[str, ...],
) -> None:
    sheet.cell(1, 1, "PotatoForge activation metric comparison")
    sheet.cell(1, 1).font = Font(bold=True, size=14, color="1F4E78")
    metadata = (
        ("Source model", str(source_path)),
        ("Calibration session", cache.calibration_session_id),
        ("Calibration version", calibration.version),
        ("Source tensor count", len(source_header.tensors)),
        ("Audited layer count", len(cache.measurements)),
        ("Baseline method", baseline_method),
        (
            "Allowed methods",
            ", ".join(sorted(allowed_methods or cache.requested_methods)),
        ),
        ("Excluded prefixes", ", ".join(excluded_prefixes) or "(none)"),
        ("Excluded suffixes", ", ".join(excluded_suffixes) or "(none)"),
        ("Metric count", len(metrics)),
        (
            "Loading note",
            "The audit cache, calibration pair, and source header are loaded once.",
        ),
    )
    for row_index, (label, value) in enumerate(metadata, start=3):
        sheet.cell(row_index, 1, label).font = Font(bold=True)
        sheet.cell(row_index, 2, value)

    rows: list[tuple[object, ...]] = []
    for metric in metrics:
        summary = profiles[metric].summary
        counts = summary["selected_action_counts"]
        assert isinstance(counts, dict)
        rows.append(
            (
                metric,
                profiles[metric].optimized.profile["profile_id"],
                summary["target_size_bytes"],
                summary["estimated_final_size_bytes"],
                summary["minimum_output_bytes"],
                summary["profile_rule_count"],
                counts.get("keep", 0),
                counts.get("convrot_w4a4", 0),
                counts.get("int8_convrot", 0),
                summary["layer_count"],
            )
        )
    _write_table(
        sheet,
        (
            "metric",
            "profile_id",
            "target_bytes",
            "estimated_output_bytes",
            "minimum_output_bytes",
            "profile_rule_count",
            "keep_count",
            "convrot_w4a4_count",
            "int8_convrot_count",
            "layer_count",
        ),
        rows,
        "MetricSummaryTable",
        start_row=16,
    )
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 72


def _write_candidate_metrics(
    sheet: Worksheet,
    cache: ActivationAuditCache,
    calibration: ActivationCalibration,
    score_rows: Mapping[str, Mapping[tuple[str, str], Mapping[str, object]]],
    candidate_methods: Sequence[str],
) -> None:
    rows: list[tuple[object, ...]] = []
    for metric, metric_rows in score_rows.items():
        for (tensor_name, method), scored in metric_rows.items():
            if method not in candidate_methods:
                continue
            candidate = cache.get(tensor_name, method)
            layer = calibration.get(tensor_name)
            if candidate is None or not isinstance(layer, LayerCalibration):
                raise ValueError(
                    f"Activation comparison cache is incomplete: {tensor_name}"
                )
            rows.append(
                (
                    metric,
                    tensor_name,
                    method,
                    candidate.action,
                    scored.get("status"),
                    scored.get("available"),
                    scored.get("storage_bytes"),
                    scored.get("objective_cost"),
                    scored.get("unavailable_reason"),
                    layer.input_features,
                    layer.output_features,
                    layer.evaluation_count,
                    layer.sample_count,
                    layer.invocation_count,
                    candidate.sample_exact_sse,
                    candidate.sample_diag_sse,
                    candidate.sample_cross_term_ratio,
                    candidate.sample_unavailable_reason,
                )
            )
    _write_table(
        sheet,
        (
            "metric",
            "tensor_name",
            "method",
            "action",
            "status",
            "available",
            "storage_bytes",
            "objective_cost",
            "unavailable_reason",
            "input_features",
            "output_features",
            "evaluation_count",
            "sample_count",
            "invocation_count",
            "sample_exact_sse",
            "sample_diag_sse",
            "sample_cross_term_ratio",
            "sample_unavailable_reason",
        ),
        rows,
        "CandidateMetricsTable",
    )


def _assignment_headers(candidate_methods: Sequence[str]) -> tuple[str, ...]:
    headers = (
        "metric",
        "tensor_name",
        "profile_id",
        "profile_action",
        "profile_method",
        "selected_storage_bytes",
        "selected_objective_cost",
        "action_variant_count",
        *(f"{method}_objective_cost" for method in candidate_methods),
    )
    if {"convrot_w4a4", "int8_convrot"}.issubset(candidate_methods):
        return (
            *headers,
            "int8_convrot_benefit",
            "int8_convrot_extra_bytes",
            "int8_convrot_benefit_per_byte",
        )
    return headers


def _assignment_rows(
    cache: ActivationAuditCache,
    source_header: SourceModelHeader,
    profiles: Mapping[str, ActivationProfileResult],
    score_rows: Mapping[str, Mapping[tuple[str, str], Mapping[str, object]]],
    candidate_methods: Sequence[str],
) -> list[tuple[object, ...]]:
    rows: list[tuple[object, ...]] = []
    for metric, generated in profiles.items():
        metric_rows = score_rows[metric]
        for tensor_name in cache.tensor_names():
            descriptor = source_header.tensors.get(tensor_name)
            if descriptor is None:
                raise ValueError(
                    f"Activation comparison source is missing: {tensor_name}"
                )
            action = resolve_profile(generated.optimized.profile, tensor_name)
            selected_method = "bf16" if action == "keep" else None
            selected_storage: int | None = (
                source_bytes(descriptor) if action == "keep" else None
            )
            selected_cost: float | None = 0.0 if action == "keep" else None
            for method in candidate_methods:
                candidate = cache.get(tensor_name, method)
                if candidate is not None and candidate.action == action:
                    selected_method = method
                    selected_storage = candidate.storage_bytes
                    scored = metric_rows.get((tensor_name, method))
                    selected_cost = (
                        None
                        if scored is None
                        else scored.get("objective_cost")
                    )
                    break
            variants = {
                cache.get(tensor_name, method).action
                for method in candidate_methods
                if cache.get(tensor_name, method) is not None
            }
            cost_by_method = {
                method: metric_rows.get((tensor_name, method), {}).get(
                    "objective_cost"
                )
                for method in candidate_methods
            }
            storage_by_method = {
                method: cache.get(tensor_name, method).storage_bytes
                for method in candidate_methods
                if cache.get(tensor_name, method) is not None
            }
            benefit = None
            extra_bytes = None
            benefit_per_byte = None
            w4a4_cost = cost_by_method.get("convrot_w4a4")
            int8_cost = cost_by_method.get("int8_convrot")
            w4a4_storage = storage_by_method.get("convrot_w4a4")
            int8_storage = storage_by_method.get("int8_convrot")
            if (
                isinstance(w4a4_cost, (int, float))
                and isinstance(int8_cost, (int, float))
                and isinstance(w4a4_storage, int)
                and isinstance(int8_storage, int)
            ):
                benefit = float(w4a4_cost) - float(int8_cost)
                extra_bytes = int8_storage - w4a4_storage
                if extra_bytes > 0:
                    benefit_per_byte = benefit / extra_bytes
            extra_values: tuple[object, ...] = (
                (benefit, extra_bytes, benefit_per_byte)
                if {"convrot_w4a4", "int8_convrot"}.issubset(candidate_methods)
                else ()
            )
            rows.append(
                (
                    metric,
                    tensor_name,
                    generated.optimized.profile["profile_id"],
                    action,
                    selected_method,
                    selected_storage,
                    selected_cost,
                    len(variants),
                    *(
                        None
                        if metric_rows.get((tensor_name, method)) is None
                        else metric_rows[(tensor_name, method)].get("objective_cost")
                        for method in candidate_methods
                    ),
                    *extra_values,
                )
            )
    return rows


def _write_profile_comparison(
    sheet: Worksheet,
    cache: ActivationAuditCache,
    profiles: Mapping[str, ActivationProfileResult],
) -> None:
    metrics = tuple(profiles)
    rows = []
    for tensor_name in cache.tensor_names():
        actions = [
            resolve_profile(profiles[metric].optimized.profile, tensor_name)
            for metric in metrics
        ]
        rows.append((tensor_name, len(set(actions)), *actions))
    _write_table(
        sheet,
        ("tensor_name", "action_variant_count", *metrics),
        rows,
        "ProfileComparisonTable",
    )
    for row in sheet.iter_rows(min_row=2, max_col=2 + len(metrics)):
        for cell in row[2:]:
            fill = _ACTION_FILLS.get(cell.value)
            if fill is not None:
                cell.fill = fill


def _write_rankings(
    sheet: Worksheet,
    cache: ActivationAuditCache,
    score_rows: Mapping[str, Mapping[tuple[str, str], Mapping[str, object]]],
    candidate_methods: Sequence[str],
    top_n: int,
) -> None:
    rows: list[tuple[object, ...]] = []
    for metric, metric_rows in score_rows.items():
        for method in candidate_methods:
            ranked = sorted(
                (
                    (tensor_name, scored)
                    for (tensor_name, candidate_method), scored in metric_rows.items()
                    if candidate_method == method
                    and isinstance(scored.get("objective_cost"), (int, float))
                ),
                key=lambda item: (-float(item[1]["objective_cost"]), item[0]),
            )
            for rank, (tensor_name, scored) in enumerate(ranked[:top_n], start=1):
                candidate = cache.get(tensor_name, method)
                rows.append(
                    (
                        metric,
                        method,
                        rank,
                        tensor_name,
                        None if candidate is None else candidate.action,
                        scored.get("objective_cost"),
                        scored.get("storage_bytes"),
                    )
                )
    _write_table(
        sheet,
        (
            "metric",
            "method",
            "rank",
            "tensor_name",
            "action",
            "objective_cost",
            "storage_bytes",
        ),
        rows,
        "RankingsTable",
    )


def _score_rows(
    report: Mapping[str, object],
) -> dict[tuple[str, str], Mapping[str, object]]:
    raw_results = report.get("results")
    if not isinstance(raw_results, list):
        raise ValueError("Activation score report results must be a list.")
    rows: dict[tuple[str, str], Mapping[str, object]] = {}
    for raw_result in raw_results:
        if not isinstance(raw_result, dict):
            raise ValueError("Activation score report contains an invalid result.")
        tensor_name = raw_result.get("tensor_name")
        method = raw_result.get("method")
        if not isinstance(tensor_name, str) or not isinstance(method, str):
            raise ValueError("Activation score result names must be strings.")
        key = (tensor_name, method)
        if key in rows:
            raise ValueError(
                f"Activation score report contains duplicate result: {tensor_name} {method}"
            )
        rows[key] = raw_result
    return rows


def _write_table(
    sheet: Worksheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    table_name: str,
    *,
    start_row: int = 1,
) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(start_row, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row_offset, values in enumerate(rows, start=1):
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(start_row + row_offset, column, value)
            if headers[column - 1] in {
                "objective_cost",
                "selected_objective_cost",
                "sample_exact_sse",
                "sample_diag_sse",
                "sample_cross_term_ratio",
                "int8_convrot_benefit",
                "int8_convrot_benefit_per_byte",
            } or headers[column - 1].endswith("_objective_cost"):
                cell.number_format = _NUMBER_FORMAT
    last_row = max(start_row + len(rows), start_row)
    last_column = len(headers)
    table = Table(
        displayName=table_name,
        ref=(
            f"A{start_row}:{get_column_letter(last_column)}{last_row}"
        ),
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = f"A{start_row + 1}"
    _set_column_widths(sheet, start_row, last_column)


def _set_column_widths(sheet: Worksheet, start_row: int, last_column: int) -> None:
    for column_index in range(1, last_column + 1):
        values = (
            len(str(cell.value))
            for column_cells in sheet.iter_cols(
                min_col=column_index,
                max_col=column_index,
                min_row=start_row,
            )
            for cell in column_cells
            if cell.value is not None
        )
        width = min(max(max(values, default=10) + 2, 12), 60)
        sheet.column_dimensions[get_column_letter(column_index)].width = width
