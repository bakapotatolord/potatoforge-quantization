import json
import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

import torch
from openpyxl import load_workbook
from safetensors.torch import save_file

from potatoforge.audits.analysis import (
    print_activation_ranking,
    print_tensor_analysis,
    write_analysis_workbook,
)
from potatoforge.audits.weight_audit import audit_bf16_source
from potatoforge.calibration import ActivationCalibration
from potatoforge.headers.source_header import read_source_model_header


class TestActivationReporting(unittest.TestCase):
    def _write_source(self, root: Path) -> Path:
        source_path = root / "source.safetensors"
        save_file(
            {
                "first.weight": torch.linspace(
                    -0.5,
                    0.5,
                    steps=512,
                    dtype=torch.bfloat16,
                ).reshape(2, 256),
                "second.weight": torch.linspace(
                    -1.0,
                    1.0,
                    steps=512,
                    dtype=torch.bfloat16,
                ).reshape(2, 256),
            },
            str(source_path),
        )
        return source_path

    def _write_calibration(self, root: Path) -> Path:
        metadata_path = root / "calibration.json"
        metadata_path.write_text(
            json.dumps(
                {
                    "format": "potatoforge_activation_calibration",
                    "version": 1,
                    "baseline_label": "bf16",
                    "activation_basis": "logical_linear_input",
                    "activation_axis": "last_dimension",
                    "layer_count": 2,
                    "layers": {
                        "first.weight": {
                            "input_features": 256,
                            "sample_count": 10,
                            "invocation_count": 2,
                            "stats_key": "first.weight.sum_x2",
                        },
                        "second.weight": {
                            "input_features": 256,
                            "sample_count": 20,
                            "invocation_count": 4,
                            "stats_key": "second.weight.sum_x2",
                        },
                    },
                }
            ),
            encoding="utf-8",
        )
        save_file(
            {
                "first.weight.sum_x2": torch.ones(256),
                "second.weight.sum_x2": torch.full((256,), 2.0),
            },
            str(root / "calibration.safetensors"),
        )
        return metadata_path

    def test_workbook_adds_activation_columns_only_when_enabled(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source_path = self._write_source(root)
            calibration = ActivationCalibration.load(
                self._write_calibration(root)
            )
            audit = audit_bf16_source(
                source_path,
                activation_calibration=calibration,
            )
            output_path = root / "analysis.xlsx"

            write_analysis_workbook(
                audit,
                read_source_model_header(source_path),
                output_path,
            )
            errors = load_workbook(output_path, data_only=True)["Errors"]

        self.assertEqual(errors.max_column, 14)
        self.assertEqual(errors["H1"].value, "Activation Error — ConvRot W4A4")
        self.assertEqual(errors["I1"].value, "Activation Samples")
        self.assertEqual(errors["J1"].value, "Activation Invocations")
        self.assertEqual(errors["K1"].value, "Activation Status")
        self.assertEqual(errors["L1"].value, "Activation Baseline")
        self.assertEqual(errors["M1"].value, "Activation Reference")
        self.assertEqual(errors["N1"].value, "Activation Candidate")
        first = audit["results"][0]["activation"]
        self.assertAlmostEqual(errors["H2"].value, first["error"])
        self.assertEqual(errors["I2"].value, first["sample_count"])
        self.assertEqual(errors["J2"].value, first["invocation_count"])
        self.assertEqual(errors["K2"].value, "ok")
        self.assertEqual(errors["L2"].value, "bf16")
        self.assertEqual(errors["M2"].value, "bf16")
        self.assertEqual(errors["N2"].value, "convrot_w4a4")

    def test_console_report_and_ranking_include_activation_data(self) -> None:
        with TemporaryDirectory() as directory:
            root = Path(directory)
            source_path = self._write_source(root)
            calibration = ActivationCalibration.load(
                self._write_calibration(root)
            )
            audit = audit_bf16_source(
                source_path,
                activation_calibration=calibration,
            )
            source_header = read_source_model_header(source_path)

            tensor_output = StringIO()
            with redirect_stdout(tensor_output):
                print_tensor_analysis(
                    audit,
                    source_header,
                    "first.weight",
                )

            audit["results"][0]["activation"]["error"] = 1.0
            audit["results"][1]["activation"]["error"] = 3.0
            ranking_output = StringIO()
            with redirect_stdout(ranking_output):
                print_activation_ranking(audit)

        rendered = tensor_output.getvalue()
        self.assertIn("Activation calibration", rendered)
        self.assertIn("Activation-aware error", rendered)
        self.assertIn("Reference:          bf16", rendered)
        self.assertIn("Status:              ok", rendered)
        self.assertIn("Activation energy:", rendered)

        ranking = ranking_output.getvalue()
        self.assertIn("Activation-aware ConvRot W4A4 ranking", ranking)
        self.assertLess(ranking.index("second.weight"), ranking.index("first.weight"))
        self.assertIn("Coverage: scored: 2", ranking)


if __name__ == "__main__":
    unittest.main()
