import unittest
from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

import torch
from openpyxl import load_workbook
from safetensors.torch import save_file

from potatoforge.audits.analysis import (
    print_tensor_analysis,
    write_analysis_workbook,
)
from potatoforge.audits.profile_optimizer import (
    estimate_profile_bytes,
    generate_profile_sweep,
    optimize_target_size,
)
from potatoforge.audits.weight_audit import audit_bf16_source
from potatoforge.headers.source_header import read_source_model_header


class TestAnalysis(unittest.TestCase):
    def _write_source(self, directory: str) -> Path:
        source_path = Path(directory) / "source.safetensors"
        save_file(
            {
                "blocks.0.attn.wq.weight": torch.ones(
                    (2, 256),
                    dtype=torch.bfloat16,
                ),
                "blocks.0.mlp.down.weight": torch.ones(
                    (2, 128),
                    dtype=torch.bfloat16,
                ),
                "embedding": torch.ones(
                    (2, 4),
                    dtype=torch.bfloat16,
                ),
            },
            str(source_path),
        )
        return source_path

    def test_workbook_reflects_audit_and_source_sizes(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)
            output_path = Path(directory) / "analysis.xlsx"

            write_analysis_workbook(audit, source_header, output_path)

            workbook = load_workbook(output_path, data_only=True)
            summary = workbook["Summary"]
            errors = workbook["Errors"]

            first_result = audit["results"][0]
            first_name = first_result["tensor_name"]
            first_offsets = source_header.tensors[first_name]["data_offsets"]
            first_bytes = first_offsets[1] - first_offsets[0]
            total_source_bytes = sum(
                descriptor["data_offsets"][1]
                - descriptor["data_offsets"][0]
                for descriptor in source_header.tensors.values()
            )

            self.assertEqual(
                summary["B3"].value,
                source_path.stat().st_size / 1024**2,
            )
            self.assertEqual(summary["B2"].value, source_path.name)
            self.assertEqual(summary["B4"].value, 3)
            self.assertEqual(summary["B5"].value, 2)
            self.assertEqual(summary["B6"].value, 2 * 256 + 2 * 128)
            self.assertEqual(summary["C10"].value, 512)
            self.assertEqual(summary["D10"].value, first_bytes / 1024**2)
            self.assertEqual(
                summary["E10"].value,
                first_bytes / total_source_bytes,
            )

            self.assertEqual(errors["A2"].value, first_name)
            self.assertEqual(errors["G1"].value, "INT8 ConvRot")
            self.assertEqual(
                errors["C1"].value,
                "W4A4 ConvRot + MSE",
            )
            self.assertEqual(
                errors["B2"].value,
                first_result["methods"]["convrot_w4a4"]["relative_l2_error"],
            )
            self.assertEqual(
                errors["C2"].value,
                first_result["methods"]["convrot_w4a4_mse"][
                    "relative_l2_error"
                ],
            )
            self.assertEqual(
                errors["D2"].value,
                first_result["methods"]["int6"]["relative_l2_error"],
            )
            self.assertEqual(
                errors["E2"].value,
                first_result["methods"]["int6_convrot"][
                    "relative_l2_error"
                ],
            )
            self.assertEqual(
                errors["F2"].value,
                first_result["methods"]["int8"]["relative_l2_error"],
            )
            self.assertEqual(
                errors["G2"].value,
                first_result["methods"]["int8_convrot"]["relative_l2_error"],
            )
            self.assertEqual(errors.max_column, 7)
            self.assertIsNone(errors["B3"].value)
            self.assertIsNone(errors["C3"].value)
            self.assertIsNone(errors["E3"].value)
            self.assertIsNone(errors["G3"].value)
            for row in range(2, 4):
                values = [
                    errors.cell(row, column).value
                    for column in range(2, 8)
                ]
                least_error = min(value for value in values if value is not None)
                for column, value in zip(range(2, 8), values):
                    self.assertEqual(
                        errors.cell(row, column).font.bold,
                        value == least_error,
                    )

            self.assertEqual(
                summary.auto_filter.ref,
                "A9:E11",
            )
            self.assertEqual(errors.freeze_panes, "B2")
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Summary",
                    "Recommendations",
                    "Errors",
                    "Trade-offs",
                    "Profiles",
                    "Profile Data",
                ],
            )

    def test_recommendations_use_the_existing_optimizer_profile(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)
            target_bytes = estimate_profile_bytes(
                source_header,
                {
                    "profile_id": "analysis",
                    "default": "keep",
                    "rules": tuple(
                        {
                            "action": "int8",
                            "prefix": result["tensor_name"],
                            "suffixes": ("",),
                        }
                        for result in audit["results"]
                    ),
                },
            )
            optimized = optimize_target_size(
                audit,
                "analysis",
                target_bytes,
                frozenset(("int8",)),
            )
            output_path = Path(directory) / "analysis.xlsx"

            write_analysis_workbook(
                audit,
                source_header,
                output_path,
                optimized,
            )

            workbook = load_workbook(
                output_path,
                data_only=False,
            )
            recommendations = workbook["Recommendations"]
            profiles = workbook["Profiles"]
            profile_data = workbook["Profile Data"]

        self.assertEqual(recommendations["B2"].value, profiles["A2"].value)
        self.assertIn("Profiles", recommendations["B3"].value)
        self.assertIn("Profile Data", recommendations["B16"].value)
        self.assertAlmostEqual(
            profile_data["E2"].value,
            audit["results"][0]["methods"]["int8"]["storage_bytes"] / 1024**2,
        )
        self.assertAlmostEqual(
            profile_data["F2"].value,
            audit["results"][0]["methods"]["int8"]["relative_l2_error"],
        )
        total_weight_energy = sum(
            result["weight_l2_sq"] for result in audit["results"]
        )
        profile_data_headers = {
            cell.value: cell.column for cell in profile_data[1]
        }
        energy_fraction_column = profile_data_headers["Energy Fraction"]
        normalized_sse_column = profile_data_headers[
            "Normalized Reconstruction SSE"
        ]
        first_result = audit["results"][0]
        first_error = first_result["methods"]["int8"]["relative_l2_error"]
        first_sse = first_error**2 * first_result["weight_l2_sq"]
        self.assertAlmostEqual(
            profile_data.cell(2, energy_fraction_column).value,
            first_result["weight_l2_sq"] / total_weight_energy,
        )
        self.assertAlmostEqual(
            profile_data.cell(2, normalized_sse_column).value,
            first_sse / total_weight_energy,
        )
        self.assertEqual(recommendations["A13"].value, "Normalized Reconstruction SSE")
        profiles_headers = {cell.value: cell.column for cell in profiles[1]}
        normalized_profile_sse = profiles.cell(
            2,
            profiles_headers["Normalized Reconstruction SSE"],
        ).value
        self.assertAlmostEqual(
            normalized_profile_sse,
            profiles["J2"].value / total_weight_energy,
        )
        self.assertEqual(recommendations["J15"].value, "Energy Fraction")
        self.assertEqual(
            recommendations["K15"].value,
            "Normalized Reconstruction SSE",
        )
        self.assertAlmostEqual(
            profiles["B2"].value,
            target_bytes / 1024**3,
        )
        self.assertEqual(recommendations.data_validations.count, 1)

    def test_profile_sweep_populates_tradeoff_chart_inputs(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)
            output_path = Path(directory) / "analysis.xlsx"
            sweep = generate_profile_sweep(
                audit,
                "analysis",
                10**9,
                frozenset(("int8",)),
            )

            write_analysis_workbook(
                audit,
                source_header,
                output_path,
                sweep[-1],
                profile_sweep=sweep,
                allowed_methods=frozenset(("int8",)),
            )

            workbook = load_workbook(output_path, data_only=False)

        tradeoffs = workbook["Trade-offs"]
        self.assertEqual(len(workbook["Profiles"]["A"]) - 1, len(sweep))
        self.assertEqual(len(tradeoffs._charts), 1)
        self.assertIn("Profiles", tradeoffs["A25"].value)
        self.assertIn("Recommendations", tradeoffs["E25"].value)

    def test_no_target_size_keeps_recommendations_explicitly_disabled(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)
            output_path = Path(directory) / "analysis.xlsx"

            write_analysis_workbook(audit, source_header, output_path)

            recommendations = load_workbook(
                output_path,
                data_only=True,
            )["Recommendations"]

        self.assertIn("No profile sweep", recommendations["B2"].value)

    def test_single_tensor_output_uses_existing_measurements(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)
            tensor_name = audit["results"][0]["tensor_name"]
            output = StringIO()

            with redirect_stdout(output):
                print_tensor_analysis(audit, source_header, tensor_name)

        rendered = output.getvalue()
        self.assertIn(tensor_name, rendered)
        self.assertIn("Shape:         2 x 256", rendered)
        self.assertIn("W4A4 ConvRot", rendered)
        self.assertIn("W4A4 ConvRot + MSE", rendered)
        self.assertNotIn("MSE Scale", rendered)
        self.assertNotIn("Recommendation", rendered)

    def test_unknown_tensor_fails_without_fuzzy_matching(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)

            with self.assertRaisesRegex(ValueError, "Tensor not found"):
                print_tensor_analysis(
                    audit,
                    source_header,
                    "blocks.0.attn.wq",
                )

    def test_workbook_refuses_overwrite_by_default(self) -> None:
        with TemporaryDirectory() as directory:
            source_path = self._write_source(directory)
            audit = audit_bf16_source(source_path)
            source_header = read_source_model_header(source_path)
            output_path = Path(directory) / "analysis.xlsx"
            output_path.write_bytes(b"existing")

            with self.assertRaises(FileExistsError):
                write_analysis_workbook(audit, source_header, output_path)


if __name__ == "__main__":
    unittest.main()
